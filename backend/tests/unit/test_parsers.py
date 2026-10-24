import io
import json
import zipfile

import docx
import pytest
from pypdf import PdfWriter

from app.ingestion.parsers import (
    NoExtractableTextError,
    ParseError,
    UnsupportedFileTypeError,
    parse_document,
)
from tests.fixtures import make_pdf


def parse(content: bytes, filename: str, mime: str | None = None):  # type: ignore[no-untyped-def]
    return parse_document(content, filename=filename, mime_type=mime)


def test_markdown_keeps_heading_path_and_title() -> None:
    text = "# Handbook\n\nIntro.\n\n## Travel\n\nFly economy.\n\n### Exceptions\n\nLong-haul.\n"
    parsed = parse(text.encode(), "handbook.md")
    assert parsed.title == "Handbook"
    assert [(s.heading, s.text) for s in parsed.sections] == [
        ("Handbook", "Intro."),
        ("Handbook > Travel", "Fly economy."),
        ("Handbook > Travel > Exceptions", "Long-haul."),
    ]


def test_markdown_ignores_headings_inside_code_blocks() -> None:
    parsed = parse(b"# Doc\n\n```\n# not a heading\n```\n", "doc.md")
    assert len(parsed.sections) == 1
    assert "# not a heading" in parsed.sections[0].text


def test_html_drops_scripts_and_keeps_sections() -> None:
    html = b"""<html><head><title>Policy</title><script>alert('x')</script></head>
    <body><nav>menu</nav><h1>Leave</h1><p>Twenty days.</p><h2>Sick</h2><p>Ten days.</p>
    <style>p{}</style></body></html>"""
    parsed = parse(html, "policy.html")
    assert parsed.title == "Policy"
    joined = " ".join(s.text for s in parsed.sections)
    assert "alert" not in joined and "menu" not in joined
    assert [(s.heading, s.text) for s in parsed.sections] == [
        ("Leave", "Twenty days."),
        ("Sick", "Ten days."),
    ]


def test_csv_rows_become_labelled_lines() -> None:
    parsed = parse(b"region,manager\nEMEA,Ana\nAPAC,Raj\n", "regions.csv")
    assert (
        parsed.sections[0].text
        == "Row 1: region: EMEA; manager: Ana\nRow 2: region: APAC; manager: Raj"
    )
    assert parsed.metadata["columns"] == ["region", "manager"]


def test_json_is_flattened_with_paths() -> None:
    data = {"title": "Config", "limits": {"upload_mb": 20}, "regions": ["eu", "us"]}
    parsed = parse(json.dumps(data).encode(), "config.json")
    assert parsed.title == "Config"
    assert "limits.upload_mb: 20" in parsed.sections[0].text
    assert "regions[1]: us" in parsed.sections[0].text


def test_pdf_text_and_page_numbers() -> None:
    parsed = parse(make_pdf(["First page text.", "Second page text."]), "doc.pdf")
    assert parsed.page_count == 2
    assert [(s.page, s.text.strip()) for s in parsed.sections] == [
        (1, "First page text."),
        (2, "Second page text."),
    ]


def test_pdf_without_text_is_reported_as_needing_ocr() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(NoExtractableTextError, match="OCR"):
        parse(buffer.getvalue(), "scan.pdf")


def test_docx_headings_and_tables() -> None:
    document = docx.Document()
    document.add_heading("Security", level=1)
    document.add_paragraph("Use MFA everywhere.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Tool"
    table.rows[0].cells[1].text = "Okta"
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse(buffer.getvalue(), "sec.docx")
    assert parsed.sections[0].heading == "Security"
    assert parsed.sections[0].text == "Use MFA everywhere."
    assert parsed.sections[1].text == "Tool | Okta"


def test_docx_decompression_bomb_is_rejected() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (50 * 1024 * 1024))
    with pytest.raises(ParseError, match="unsafe size"):
        parse(buffer.getvalue(), "bomb.docx")


@pytest.mark.parametrize(
    ("content", "filename", "error"),
    [
        (b"MZ\x90\x00binary", "tool.exe", UnsupportedFileTypeError),
        (b"not a pdf at all", "fake.pdf", ParseError),
        (b"plain text", "fake.docx", ParseError),
        (b"\x00\x01\x02binary", "data.txt", ParseError),
        (b"{not json", "broken.json", ParseError),
        (b"%PDF-1.4 truncated garbage", "broken.pdf", ParseError),
        (b"   \n\n  ", "empty.txt", NoExtractableTextError),
    ],
)
def test_malformed_inputs_raise_typed_errors(
    content: bytes, filename: str, error: type[Exception]
) -> None:
    with pytest.raises(error):
        parse(content, filename)


def test_type_detected_from_mime_when_no_extension() -> None:
    assert parse(b"hello world", "upload", "text/plain").kind == "text"


def test_legacy_encoding_falls_back_to_cp1252() -> None:
    parsed = parse("café résumé".encode("cp1252"), "notes.txt")
    assert parsed.sections[0].text == "café résumé"


def test_json_key_amplification_is_capped() -> None:
    # Review finding: a 26 KB file used to expand to ~40 MB of text (key repeated per leaf).
    bomb = json.dumps({"k" * 20000: [1] * 2000}).encode()
    parsed = parse(bomb, "bomb.json")  # long keys are truncated, so this stays small
    assert sum(len(s.text) for s in parsed.sections) < 4 * len(bomb) + 1_000_000


def test_json_expansion_beyond_budget_is_rejected() -> None:
    wide = json.dumps({"x" * 100: list(range(40_000))}).encode()  # ~240 KB -> >4 MB of text
    with pytest.raises(ParseError, match="too much text"):
        parse(wide, "wide.json")


def test_csv_header_amplification_is_capped() -> None:
    # 20 columns with 95-char names: every row repeats ~2 KB of header text for 20 bytes of data.
    header = ",".join(f"c{i:02d}" + "h" * 92 for i in range(20))
    rows = "\n".join(",".join("1" for _ in range(20)) for _ in range(19_000))
    with pytest.raises(ParseError, match="too much text"):
        parse(f"{header}\n{rows}".encode(), "wide.csv")
